from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path

doc = Document()
s = doc.sections[0]
s.top_margin = s.bottom_margin = Inches(.65)
s.left_margin = s.right_margin = Inches(.7)
s.page_width, s.page_height = Inches(8.5), Inches(11)
for name in ['Normal', 'Title', 'Heading 1']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing = 1.05
doc.styles['Title'].font.size = Pt(21)
doc.styles['Heading 1'].font.size = Pt(13)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(5)

def p(t): doc.add_paragraph(t)
def h(t): doc.add_heading(t,1)
def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers): c.text=txt
    rep=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row): c.text=txt
    for i,row in enumerate(t.rows):
        for c,w in zip(row.cells,widths):
            c.width=Inches(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'75'); e.set(qn('w:type'),'dxa'); mar.append(e)
            pr.append(mar)
            if i==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EAEAEA'); pr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(2)
                para.paragraph_format.line_spacing=1
                for run in para.runs: run.font.size=Pt(9.5); run.bold=(i==0)

doc.add_paragraph('Plan de pruebas de Crosswalker', 'Title')
p('Avance preliminar • Versión 0.1 • 8 de octubre de 2026')
h('1 Introducción')
p('Crosswalker es un juego de rompecabezas para Android que consiste en ayudar a peatones a cruzar la calle de forma segura. El jugador elige cruces, asigna acciones de seguridad y coloca protecciones con un presupuesto de fichas. Incluye 12 niveles en tres mundos y funciona sin conexión. Este plan propone una revisión básica antes de su entrega.')
h('2 Objetivos')
p('Comprobar que las reglas del juego funcionen correctamente, que el avance se guarde y que los controles sean fáciles de entender. Detectar errores que impidan completar una partida o desbloquear niveles.')
h('3 Alcance')
p('Se revisarán el inicio del juego, la selección de niveles, la asignación de acciones, el límite de fichas, la simulación del cruce, los resultados, las estrellas y el guardado del progreso. Se incluyen validación de requerimientos, prueba alfa, regresión, pruebas dinámicas, funcionales, de usabilidad y no funcionales.')
h('4 Exclusiones')
p('No se probarán versiones para iOS o web, funciones multijugador ni servicios en línea. Tampoco se cubrirán todos los modelos de teléfono ni una auditoría completa de seguridad.')
h('5 Casos preliminares de validación y pruebas funcionales')
p('Cada fila representa un requerimiento básico y su caso de prueba de caja negra. Todos los casos están pendientes de ejecución.')
table(['Caso', 'Acción de prueba', 'Resultado esperado'],[
('CP01 Inicio','Abrir el juego y entrar al primer nivel.','El nivel carga y permite preparar el cruce.'),
('CP02 Plan','Elegir un cruce y asignar una acción de seguridad.','La selección queda reflejada en el plan.'),
('CP03 Fichas','Intentar superar el presupuesto de fichas.','El juego impide exceder el presupuesto.'),
('CP04 Cruce','Ejecutar un plan seguro y otro que provoque un choque.','El primero permite ganar; el segundo permite reintentar.'),
('CP05 Progreso','Ganar un nivel, cerrar y abrir la aplicación.','Conserva las estrellas y desbloquea el siguiente nivel.'),
('CP06 Estrellas','Repetir un nivel con una puntuación menor.','Se conserva la mejor puntuación obtenida.')
],[1.05,3.0,3.05])
doc.add_page_break()
h('6 Responsables y calendario')
p('Fechas propuestas, sujetas a confirmación. Se asignarán nombres a los roles de desarrollo y pruebas antes de iniciar. Ninguna prueba se considera aprobada todavía.')
table(['Prueba', 'Actividad', 'Responsable', 'Fecha tentativa'],[
('Validación de requerimientos','Ejecutar CP01 a CP06 y comparar con las reglas.','Pruebas','12/10/2026'),
('Alfa','Evaluación heurística de claridad, controles, mensajes y facilidad para reintentar.','Pruebas','13/10/2026'),
('Dinámicas','Caja blanca: revisar rutas de presupuesto, colisiones, emergencia y cálculo de estrellas.','Desarrollo','14/10/2026'),
('Funcionales','Caja negra: ejecutar CP01 a CP06 desde la interfaz.','Pruebas','15/10/2026'),
('Regresión','Repetir los casos después de corregir errores. Pendiente de una versión de prueba estable.','Desarrollo y pruebas','16/10/2026'),
('Usabilidad','Observar a tres usuarios al preparar y ejecutar un cruce. Pendiente de una versión de prueba estable.','Pruebas y usuarios','19/10/2026'),
('Rendimiento','Medir carga del nivel y observar una partida de 10 minutos.','Desarrollo','20/10/2026'),
('Compatibilidad','Repetir una partida en Android 7.0 y en una versión posterior, sin conexión.','Pruebas','20/10/2026')
],[1.15,3.15,1.3,1.5])
h('7 Necesidades ambientales')
p('Una computadora con Android Studio, el proyecto del juego y un emulador o teléfono con Android 7.0 o superior. Para compatibilidad se usará también una versión posterior. Se requiere una instalación limpia, otra con progreso guardado y una hoja para registrar resultados y capturas de errores.')
h('8 Criterios de aceptación')
p('Los seis casos básicos deben aprobarse y no debe haber errores que bloqueen partidas o borren el progreso. Como metas preliminares, cada nivel debe cargar en menos de 5 segundos y el juego debe funcionar durante 10 minutos sin cierres ni bloqueos en los equipos de prueba. Debe funcionar sin conexión en ambos entornos Android. Al menos dos de tres usuarios deben preparar y ejecutar un cruce sin ayuda externa, después de leer las indicaciones del juego.')
h('9 Aprobaciones')
p('El responsable de pruebas registrará los resultados y el responsable de desarrollo corregirá los errores críticos. La persona encargada del proyecto revisará la evidencia y aprobará la versión cuando cumpla los criterios anteriores.')
p('Aprobado por: ____________________    Fecha: ______________')
out=Path('outputs/Plan_de_pruebas_Crosswalker.docx')
out.parent.mkdir(exist_ok=True)
doc.save(out)
print(out.resolve())
